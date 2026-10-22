"""Импорт выписок бирж: CSV и Excel, с авто-разбором колонок.

Зачем отдельный модуль. Раньше выписка принималась только в CSV и только в
трёх жёстко зашитых форматах (Bybit/OKX/внутренний). На практике выписки
приходят в виде .xlsx, с произвольным порядком колонок и русскими
заголовками («Дата», «Торговая пара», «Тип операции», «Количество»,
«Цена», «Сумма»), плюс отдельными блоками вводов и выводов.

Что делает модуль:
  * читает CSV (разделитель определяется автоматически) и XLSX (openpyxl);
  * по заголовкам находит колонки: время, пара, сторона/тип, количество,
    цена, сумма, комиссия, валюта;
  * определяет биржу (Bybit / OKX / Binance / универсальный);
  * приводит строки к общему ParsedRow — тот же формат, что у API-синка и
    CSV-импорта, поэтому запись идёт через persist_cex_rows и дедупликация
    работает одинаково.

ВАЖНО про валюту. Цена в выписке — это quote-валюта пары (USDT, BTC и т.п.),
а не рубли. Поля fiat_value_rub / cost_basis_rub во всём проекте — РУБЛИ.
Поэтому парсер отдаёт quote_asset/quote_amount, а перевод в рубли делает
price_oracle по курсу на дату операции. Если котировки нет — сделка пишется
без рублевой оценки и попадает в отчёт как «нет котировки», а не как
неверное число.

Вводы и выводы (пополнение/вывод с биржи) — это ПЕРЕМЕЩЕНИЕ активов, а не
доход и не реализация: они не должны попадать в налоговую базу. Им соответствуют
TRANSFER_IN / TRANSFER_OUT, которые лотов не открывают и не гасят.
"""

from __future__ import annotations

import csv
import io
import logging
import re
import time
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Iterable, Optional

from utils.crypto_db import TxType
from utils.csv_import import ParsedRow, _parse_internal
from utils.price_oracle import normalize_iso_date

logger = logging.getLogger(__name__)


# ----------------------------------------------------- распознавание колонок

# Синонимы заголовков. Ключ — семантическая роль, значение — возможные
# названия колонок (в нижнем регистре, без учёта регистра букв).
COLUMN_SYNONYMS: dict[str, tuple[str, ...]] = {
    "time": (
        "utctime", "utc time", "time", "timestamp", "date", "datetime",
        "date(utc)", "ctime", "created time", "trade time", "time(utc)",
        "дата", "время", "дата и время", "дата операции", "дата/время",
        "время операции", "дата исполнения", "дата сделки",
    ),
    "pair": (
        "symbol", "pair", "market", "instid", "instrument", "trading pair",
        "symbol(pair)", "торговая пара", "пара", "инструмент", "тикер",
    ),
    "side": (
        "side", "type", "operation", "action", "operation type", "trade type",
        "buy/sell", "direction", "тип", "операция", "действие", "тип операции",
        "направление", "вид операции", "buy_sell",
    ),
    "amount": (
        "qty", "quantity", "amount", "base amount", "executed amount",
        "execution amount", "volume", "vol", "size", "filled amount",
        "fillsz", "fill size", "filled size", "sz", "filled qty", "exec qty",
        "количество", "объем", "объём", "количество (base)", "amount (base)",
        "исполнено", "количество (base asset)",
    ),
    "price": (
        "price", "exec price", "execution price", "avg price", "average price",
        "fill price", "fillpx", "fill px", "px", "avgprice", "limit price",
        "цена", "цена исполнения", "средняя цена",
    ),
    "quote_amount": (
        "total", "quote amount", "quote qty", "quote quantity", "value",
        "turnover", "traded value", "сумма", "общая сумма", "сумма сделки",
        "итого", "сумма в валюте", "объём в валюте",
    ),
    "quote_asset": (
        "quote asset", "quote currency", "quote symbol", "payment currency",
        "quote coin", "котируемая валюта", "валюта пары", "валюта котировки",
    ),
    "fee": (
        "fee", "commission", "trading fee", "fee amount", "комиссия",
        "комиссия (base)", "комиссия (quote)",
    ),
    "fee_asset": (
        "fee coin", "fee currency", "fee asset", "fee symbol",
        "комиссия (валюта)", "валюта комиссии",
    ),
    "asset": (
        "coin", "currency", "asset", "token", "symbol(asset)", "ticker(asset)",
        "валюта", "актив", "токен", "монета", "тикер актива",
    ),
    "status": ("status", "state", "статус", "состояние"),
    "address": ("address", "addr", "адрес", "адрес кошелька", "wallet"),
    "txid": (
        "txid", "tx id", "tx hash", "transaction id", "transaction hash",
        "hash", "order id", "orderid", "идентификатор транзакции", "id",
    ),
}

# Значения колонки «тип/сторона» → тип операции в справочнике TxType.
SIDE_TYPES: dict[str, str] = {
    "buy": TxType.BUY,
    "b": TxType.BUY,
    "покупка": TxType.BUY,
    "покупка (маркет)": TxType.BUY,
    "купить": TxType.BUY,
    "sell": TxType.SELL,
    "s": TxType.SELL,
    "продажа": TxType.SELL,
    "продать": TxType.SELL,
    "deposit": TxType.TRANSFER_IN,
    "пополнение": TxType.TRANSFER_IN,
    "ввод": TxType.TRANSFER_IN,
    "зачисление": TxType.TRANSFER_IN,
    "credit": TxType.TRANSFER_IN,
    "replenishment": TxType.TRANSFER_IN,
    "withdrawal": TxType.TRANSFER_OUT,
    "withdraw": TxType.TRANSFER_OUT,
    "вывод": TxType.TRANSFER_OUT,
    "вывод средств": TxType.TRANSFER_OUT,
    "debit": TxType.TRANSFER_OUT,
    "staking": TxType.STAKING,
    "стейкинг": TxType.STAKING,
    "reward": TxType.REWARD,
    "награда": TxType.REWARD,
    "earn": TxType.REWARD,
    "interest": TxType.INTEREST,
    "процент": TxType.INTEREST,
    "airdrop": TxType.AIRDROP,
    "airdrop rewards": TxType.AIRDROP,
    "аирдроп": TxType.AIRDROP,
    "income": TxType.INCOME_OTHER,
    "доход": TxType.INCOME_OTHER,
    "gift": TxType.GIFT_IN,
    "подарок": TxType.GIFT_IN,
}

# Типы, которые не меняют налоговую базу (перемещения) — по ним не нужна
# цена и не нужен расчёт прибыли.
TRANSFER_TYPES = {TxType.TRANSFER_IN, TxType.TRANSFER_OUT}

# Типы-доходы: по ним считается доход по ст. 214 НК.
INCOME_ROWS = {
    TxType.STAKING, TxType.REWARD, TxType.INTEREST, TxType.AIRDROP,
    TxType.GIFT_IN, TxType.INCOME_OTHER, TxType.REFUND,
}

# Quote-валюта пары: по ней определяется, что в строке база, а что котировка.
QUOTE_ASSETS = (
    "FDUSD", "USDT", "USDC", "TUSD", "BUSD", "DAI", "USDE", "PYUSD",
    "USD", "EUR", "RUB", "TRY", "BTC", "ETH", "BNB", "TON",
)


# Статусы, при которых операция не состоялась — в базу не идёт.
FAILED_STATUSES = (
    "cancel", "cancelled", "canceled", "fail", "failed", "reject",
    "rejected", "отмен", "отклон", "ошибк", "не исполн",
)


# ----------------------------------------------------- чтение таблицы


def _sniff_delimiter(text: str) -> str:
    """Разделитель CSV: запятая, точка с запятой или табуляция."""
    head = text[:4000]
    counts = {d: head.count(d) for d in (",", ";", "\t", "|")}
    best = max(counts, key=lambda d: counts[d])
    return best if counts[best] else ","


def _read_csv(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig", errors="replace")
    delimiter = _sniff_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [[str(c) for c in row] for row in reader]


def _read_xlsx(content: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = wb[wb.sheetnames[0]]
    out: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        cells = ["" if c is None else str(c) for c in row]
        if any(c.strip() for c in cells):
            out.append(cells)
    wb.close()
    return out


def read_table(content: bytes, filename: str = "") -> list[list[str]]:
    """Читает CSV или XLSX в список строк. Формат — по сигнатуре файла."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm", ".xls")) or content[:2] == b"PK":
        try:
            return _read_xlsx(content)
        except Exception as e:
            logger.warning(f"statement: xlsx не прочитан ({e}) — пробуем csv")
    return _read_csv(content)


# ----------------------------------------------------- сопоставление колонок


def _norm_header(value: str) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def map_columns(headers: list[str]) -> dict[str, int]:
    """Семантическая роль → индекс колонки. Возвращает только найденные."""
    norm = [_norm_header(h) for h in headers]
    found: dict[str, int] = {}
    for role, names in COLUMN_SYNONYMS.items():
        for idx, h in enumerate(norm):
            if role in found:
                break
            if not h:
                continue
            if h in names:
                found[role] = idx
                break
        if role in found:
            continue
        # Совпадение по началу заголовка — «Amount (BTC)», «Цена, USDT».
        # Именно по началу, а не «внутри»: иначе «Order Type» (лимит/маркет)
        # перебил бы колонку «Side» (покупка/продажа) и сделки перестали бы
        # делиться на покупки и продажи.
        for idx, h in enumerate(norm):
            if not h:
                continue
            for name in names:
                if len(name) >= 4 and h.startswith(name):
                    found[role] = idx
                    break
            if role in found:
                break
    return found


def detect_format(headers: list[str]) -> str:
    """По заголовкам — формат выписки."""
    h = {_norm_header(x) for x in headers}
    cols = map_columns(headers)
    if "utctime" in h and "symbol" in h:
        return "bybit"
    if "ctime" in h and "instid" in h:
        return "okx"
    if "pair" in h and ("date(utc)" in h or "utc time" in h or "time" in h):
        return "binance"
    if {"ts", "asset", "type"} <= h:
        return "internal"
    if "time" in cols and ("pair" in cols or "asset" in cols) and "side" in cols:
        return "generic"
    if "time" in cols and ("pair" in cols or "asset" in cols):
        return "generic"
    return "unknown"


# ----------------------------------------------------- разбор значений


def _decimal_str(value: str) -> str:
    """Пустое/нечисловое → ''. Иначе нормализованная строка числа."""
    s = str(value or "").strip().replace(" ", "").replace(" ", "")
    if not s:
        return ""
    s = s.replace(",", ".")
    if s.count(".") > 1:      # «1.234.567,89» → «1234567.89»
        head, _, tail = s.rpartition(".")
        s = head.replace(".", "") + "." + tail
    try:
        Decimal(s)
    except (InvalidOperation, ValueError):
        return ""
    return s


def _safe_mul(a: str, b: str) -> str:
    """Произведение двух чисел-строк через Decimal (без артефактов float)."""
    try:
        result = Decimal(a) * Decimal(b)
    except (InvalidOperation, TypeError, ValueError):
        return "0"
    if result == result.to_integral_value():
        return str(result.quantize(Decimal("1")))
    return str(result.normalize())


def split_pair(pair: str) -> tuple[str, str]:
    """«BTCUSDT» / «BTC-USDT» / «BTC/USDT» → (BTC, USDT).

    Сначала разделитель: «BTC-USDT-SWAP» — это BTC/USDT с суффиксом рынка,
    и разбор по суффиксу дал бы «BTC-» вместо «BTC».
    """
    s = str(pair or "").strip().upper().replace("/", "-").replace("_", "-")
    if not s:
        return "", ""
    parts = [p for p in s.split("-") if p]
    if len(parts) >= 2:
        return parts[0], parts[1]
    for quote in QUOTE_ASSETS:
        if s.endswith(quote) and len(s) > len(quote):
            return s[:-len(quote)], quote
    return s, ""


def _parse_ts(value: str) -> Optional[int]:
    """Время из строки выписки: ISO, «дд.мм.гггг», epoch (сек/мс)."""
    s = str(value or "").strip()
    if not s:
        return None
    if s.isdigit():
        ts = int(s)
        if ts > 10_000_000_000:   # миллисекунды
            ts //= 1000
        if ts > 1_000_000_000:
            return ts
    iso = normalize_iso_date(s)
    if iso:
        try:
            return int(datetime.strptime(iso, "%Y-%m-%d").replace(
                tzinfo=timezone.utc).timestamp())
        except ValueError:
            pass
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y %H:%M:%S",
                "%d.%m.%Y %H:%M", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d %H:%M:%S"):
        try:
            return int(time.mktime(time.strptime(s, fmt)))
        except ValueError:
            continue
    # «2025-03-01T10:00:00Z» и подобные ISO-8601.
    try:
        return int(datetime.fromisoformat(s.replace("Z", "+00:00")).timestamp())
    except ValueError:
        return None


def _side_to_type(value: str) -> str:
    """Значение колонки «тип/сторона» → тип операции (пусто, если не поняли)."""
    s = str(value or "").strip().lower()
    if not s:
        return ""
    if s in SIDE_TYPES:
        return SIDE_TYPES[s]
    for key, tx_type in SIDE_TYPES.items():
        if key in s:
            return tx_type
    return ""


# ----------------------------------------------------- основной разбор


def parse_statement(
    content: bytes,
    filename: str = "",
    source_hint: str = "",
) -> tuple[list[ParsedRow], list[str]]:
    """Разбирает выписку биржи (CSV/XLSX) в список ParsedRow.

    Возвращает (rows, errors). errors — понятные человеку сообщения.
    """
    try:
        table = read_table(content, filename)
    except Exception as e:
        logger.warning(f"statement: файл не прочитан: {e}")
        return [], [f"не удалось прочитать файл: {e}"]
    if not table:
        return [], ["пустой файл"]

    headers = table[0]
    body = table[1:]
    fmt = source_hint or detect_format(headers)
    if fmt == "internal":
        return _parse_internal(body)
    if fmt == "bybit":
        rows, errors = _parse_bybit_header(headers, body)
        if rows:
            return rows, errors
    if fmt == "okx":
        rows, errors = _parse_okx_header(headers, body)
        if rows:
            return rows, errors
    return _parse_generic(headers, body, fmt)


def _exchange_of(fmt: str, headers: list[str]) -> str:
    """Имя источника для external_id: bybit / okx / binance / csv."""
    if fmt in ("bybit", "okx", "binance"):
        return fmt
    h = {_norm_header(x) for x in headers}
    if "utctime" in h and "symbol" in h:
        return "bybit"
    if "ctime" in h and "instid" in h:
        return "okx"
    if "pair" in h:
        return "binance"
    return "csv"


def _parse_bybit_header(headers: list[str], body: Iterable[list[str]]):
    """Bybit: utcTime, symbol, side, qty, price, fee."""
    cols = map_columns(headers)
    if "time" not in cols or "pair" not in cols:
        return [], []
    return _parse_generic(headers, body, "bybit")


def _parse_okx_header(headers: list[str], body: Iterable[list[str]]):
    """OKX: cTime, instId, side, fillSz, fillPx."""
    cols = map_columns(headers)
    if "time" not in cols or "pair" not in cols:
        return [], []
    return _parse_generic(headers, body, "okx")


def _parse_generic(headers: list[str], body: Iterable[list[str]],
                   fmt: str = "generic") -> tuple[list[ParsedRow], list[str]]:
    """Универсальный разбор по найденным колонкам.

    Работает с любым экспортом, где есть «дата + пара/валюта + тип»:
    Bybit, OKX, Binance, русские заголовки, произвольный порядок колонок.
    """
    cols = map_columns(headers)
    exchange = _exchange_of(fmt, headers)
    rows: list[ParsedRow] = []
    errors: list[str] = []
    skipped = 0

    if "time" not in cols:
        return [], [f"не найдена колонка с датой (заголовки: {headers!r})"]

    def cell(row: list[str], role: str) -> str:
        idx = cols.get(role)
        if idx is None or idx >= len(row):
            return ""
        return str(row[idx] or "").strip()

    for i, r in enumerate(body, start=2):
        if not any(str(c or "").strip() for c in r):
            continue

        status = cell(r, "status").lower()
        if status and any(word in status for word in FAILED_STATUSES):
            # Отменённая/неисполненная сделка в налоговую базу не идёт.
            skipped += 1
            continue

        ts = _parse_ts(cell(r, "time"))
        if ts is None:
            errors.append(f"строка {i}: не читается дата {cell(r, 'time')!r}")
            continue

        raw_side = cell(r, "side")
        tx_type = _side_to_type(raw_side)
        pair_raw = cell(r, "pair")
        asset_col = cell(r, "asset")

        asset, quote = split_pair(pair_raw) if pair_raw else ("", "")
        if not asset:
            asset = asset_col.strip().upper()
        if not asset:
            errors.append(f"строка {i}: не определён актив (тип {raw_side!r})")
            continue

        amount = _decimal_str(cell(r, "amount"))
        price = _decimal_str(cell(r, "price"))
        quote_amount = _decimal_str(cell(r, "quote_amount"))
        quote_asset = cell(r, "quote_asset").strip().upper() or quote

        if not amount:
            errors.append(f"строка {i}: пустое количество")
            continue

        if tx_type in TRANSFER_TYPES or tx_type in INCOME_ROWS:
            # Ввод/вывод и доходы: пары нет — только актив и количество.
            quote_amount = ""
            quote_asset = ""
        else:
            if not tx_type:
                errors.append(f"строка {i}: непонятный тип операции {raw_side!r}")
                continue
            if not quote_amount and price:
                quote_amount = _safe_mul(amount, price)

        ext_id = _external_id(exchange, ts, asset, tx_type, amount, cell(r, "txid"))
        rows.append(ParsedRow(
            ts=ts, asset=asset, tx_type=tx_type,
            amount=amount, fiat_value_rub="", cost_basis_rub="",
            note=_note_for(tx_type, exchange, raw_side, pair_raw),
            external_id=ext_id,
            quote_asset=quote_asset, quote_amount=quote_amount,
            source=exchange,
        ))
    if skipped:
        errors.append(f"пропущено строк со статусом «отменено/не исполнено»: {skipped}")
    return rows, errors


def _note_for(tx_type: str, exchange: str, raw_side: str, pair_raw: str) -> str:
    label = TxType.label(tx_type)
    pair = f" {pair_raw}" if pair_raw else ""
    return f"{exchange}: {label}{pair}"


def _external_id(exchange: str, ts: int, asset: str, tx_type: str,
                 amount: str, txid: str = "") -> str:
    """Стабильный ключ дедупликации строки выписки."""
    if txid:
        return f"{exchange}:{txid}"
    return f"{exchange}:{ts}:{asset}:{tx_type}:{amount}"
